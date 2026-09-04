"""
LLM-as-judge: YAML judge definitions, label/logprob scoring, and the
per-judge scores.xlsx workbook. See README_llm.md's "Judges" section for
the full schema and worked examples.
"""

import json
import logging
import math
import os
import re
from pathlib import Path

import aiohttp
import yaml
from webdav3.client import Client
from webdav3.exceptions import RemoteResourceNotFound

from src.nextcloud.batch import _normalize_list, _sidecar_from_dict, _stt_models_from_dict
from src.nextcloud.constants import (
    JUDGE_TOP_LOGPROBS,
    JUDGES_SUBFOLDER,
    LLM_SUBFOLDER,
    SIDECAR_SUFFIXES,
    TRANSCRIPT_SUBFOLDER,
)
from src.nextcloud.results_table import _scan_prompt_cells
from src.nextcloud.webdav_utils import _download_text

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Judges (LLM-as-judge)
# ---------------------------------------------------------------------------


def _parse_judge_scale(scale) -> dict:
    """Parse a judge YAML's `scale:` field into a normalized shape.

    Accepts:
    - `"1-5"` (a string, `low-high`) or a 2-element numeric list `[1, 5]` ->
      an inclusive numeric range.
    - A >2-element numeric list, e.g. `[1, 3, 5]` -> an explicit discrete
      numeric set (supports non-contiguous scales).
    - A list of non-numeric strings, e.g. `["poor", "fair", "good"]` -> a
      label scale. Must be given in ASCENDING QUALITY ORDER — required for
      `kind: logprob`'s ordinal mapping (see `_score_from_logprobs`).

    Args:
        scale: The raw `scale:` value from a parsed judge YAML.

    Returns:
        `{"type": "numeric"|"label", "values": list}`.

    Raises:
        ValueError: If `scale` doesn't match any of the above shapes.
    """
    if isinstance(scale, str):
        m = re.match(r"^\s*(-?\d+)\s*-\s*(-?\d+)\s*$", scale)
        if m:
            low, high = int(m.group(1)), int(m.group(2))
            return {"type": "numeric", "values": list(range(low, high + 1))}
    elif isinstance(scale, list) and scale:
        try:
            nums = [int(v) if float(v).is_integer() else float(v) for v in scale]
        except (TypeError, ValueError):
            nums = None
        if nums is not None:
            if len(nums) == 2:
                low, high = nums
                if isinstance(low, int) and isinstance(high, int):
                    return {"type": "numeric", "values": list(range(low, high + 1))}
            return {"type": "numeric", "values": nums}
        return {"type": "label", "values": [str(v) for v in scale]}

    raise ValueError(
        f"Judge scale must be a numeric range/list or a list of labels, got: {scale!r}"
    )


def _normalize_anchors(anchors, scale: dict) -> dict[str, str]:
    """Normalize a judge YAML's `anchors:` field into `{scale_value_str: description}`.

    A user should never be forced to write one anchor per scale point,
    especially on a fine-grained numeric scale (e.g. 1-100). Two input
    shapes are accepted:
    - dict (explicit): `{scale_value: description}` — any sparse subset,
      used as-is at exactly those positions.
    - list (auto-distributed): plain list of description strings with no
      explicit positions — spread evenly across `scale["values"]`,
      endpoints inclusive. For N anchors over M scale values, anchor `i`
      lands at scale index `round(i * (M-1) / (N-1))` (N==1 -> the middle
      scale value).

    Args:
        anchors: The raw `anchors:` value (dict, list, or None/absent).
        scale: Parsed scale, as returned by `_parse_judge_scale`.

    Returns:
        `{str(scale_value): description}`, possibly empty.
    """
    if not anchors:
        return {}
    if isinstance(anchors, dict):
        return {str(k): str(v) for k, v in anchors.items()}

    values = scale["values"]
    n, m = len(anchors), len(values)
    if n == 1:
        positions = [m // 2]
    else:
        positions = [round(i * (m - 1) / (n - 1)) for i in range(n)]
    return {str(values[pos]): str(desc) for pos, desc in zip(positions, anchors)}


def _format_scale_section(scale: dict) -> str:
    """Render the scale portion of a judge prompt."""
    if scale["type"] == "numeric":
        return f"Scale: rate from {scale['values'][0]} to {scale['values'][-1]}."
    return f"Scale: choose one of: {', '.join(scale['values'])}."


def _format_anchors_section(anchors, scale: dict) -> str:
    """Render the anchors portion of a judge prompt, or "" if none given."""
    normalized = _normalize_anchors(anchors, scale)
    if not normalized:
        return ""
    lines = "\n".join(f"- {value}: {desc}" for value, desc in normalized.items())
    return f"Anchors:\n{lines}"


def _format_few_shot_section(few_shot) -> str:
    """Render the few-shot examples portion of a judge prompt, or "" if none given."""
    if not few_shot:
        return ""
    blocks = []
    for ex in few_shot:
        block = f"Response: {ex.get('response', '')}\nScore: {ex.get('score', '')}"
        if ex.get("rationale"):
            block += f"\nRationale: {ex['rationale']}"
        blocks.append(block)
    return "Examples:\n" + "\n\n".join(blocks)


def _format_anti_pattern_section(anti_pattern) -> str:
    """Render the anti-pattern portion of a judge prompt, or "" if none given."""
    items = _normalize_list(anti_pattern)
    if not items:
        return ""
    lines = "\n".join(f"- {item}" for item in items)
    return f"Do NOT reward:\n{lines}"


def _build_judge_template(judge_data: dict) -> str:
    """Assemble one judge's fixed-order prompt template from its parsed YAML.

    Fixed section order (only non-empty/optional ones are included):
      1. `judge_data["prompt"]` (stripped) — the user's own rubric text.
      2. Scale section — always present.
      3. Anchors section — only if `anchors` given.
      4. Few-shot examples section — only if `few_shot` given.
      5. Anti-pattern section — only if `anti_pattern` given.
      6. "Response to evaluate:\\n{response}" — always present.
      7. A kind-specific answer-format instruction — always present.

    Sections are joined with blank lines. Returns the unrendered template
    (still contains `{response}`, and `{transcript}`/`{context}` if section
    1 referenced them) — built once per judge at discovery time, not per
    scored cell.

    Args:
        judge_data: The raw parsed judge YAML dict (already validated).

    Returns:
        Unrendered judge prompt template.
    """
    scale = judge_data["_scale"]
    sections = [
        judge_data["prompt"].strip(),
        _format_scale_section(scale),
        _format_anchors_section(judge_data.get("anchors"), scale),
        _format_few_shot_section(judge_data.get("few_shot")),
        _format_anti_pattern_section(judge_data.get("anti_pattern")),
        "Response to evaluate:\n{response}",
        (
            "Respond with ONLY the label above, nothing else. No explanation, no punctuation."
            if judge_data["kind"] == "label"
            else "Respond with a single token: the number only. Do not explain your reasoning."
        ),
    ]
    return "\n\n".join(s for s in sections if s)


def _render_judge_prompt(
    template: str,
    response: str,
    transcript: str = "",
    context: str = "",
    extra_fields: dict[str, str] | None = None,
) -> str:
    """Render a judge prompt template with the response text being judged.

    Args:
        template: Unrendered template from `_build_judge_template`.
        response: The `.md` output text being judged.
        transcript: Original transcript text, substituted only if the
            judge's own text references `{transcript}`.
        context: Sidecar-provided context, substituted only if referenced.
        extra_fields: Any non-schema sidecar/batch-row fields from the
            judged file's own row (see `_sidecar_from_dict`) — e.g. a
            user-added `ground_truth` column, referenced as `{ground_truth}`
            in the judge's `prompt:` text for real reference-based judging,
            not just reference-free quality checks.

    Returns:
        Rendered judge prompt text.
    """
    rendered = (
        template.replace("{response}", response)
        .replace("{transcript}", transcript)
        .replace("{context}", context)
    )
    for key, value in (extra_fields or {}).items():
        rendered = rendered.replace("{" + key + "}", value)
    return rendered


def _parse_judge_yaml(text: str, name: str) -> dict:
    """Parse and validate one `judges/<name>.yaml`.

    Args:
        text: Raw YAML file contents.
        name: Judge name (the file's stem), used in error messages and as
            the output filename prefix.

    Returns:
        `{"name", "llm", "kind", "jurisdiction", "scale", "template"}`.

    Raises:
        ValueError: If a required field is missing, `kind` is invalid, or
            the scale doesn't parse (see `_parse_judge_scale`). Callers
            should catch this per-file so one malformed judge doesn't abort
            discovery of every other judge.
    """
    data = yaml.safe_load(text) or {}
    for field in ("prompt", "jurisdiction", "scale", "llm", "kind"):
        if not data.get(field):
            raise ValueError(f"judges/{name}.yaml: missing required field '{field}'")
    if data["kind"] not in ("label", "logprob"):
        raise ValueError(
            f"judges/{name}.yaml: kind must be 'label' or 'logprob', got {data['kind']!r}"
        )

    scale = _parse_judge_scale(data["scale"])
    data["_scale"] = scale  # stashed for _build_judge_template, not part of the public return shape
    return {
        "name": name,
        "llm": str(data["llm"]),
        "kind": data["kind"],
        "jurisdiction": _normalize_list(data["jurisdiction"]),
        "scale": scale,
        "template": _build_judge_template(data),
    }


def _collect_judges(
    client: Client, remote_roots: list[str], tmp_dir: Path
) -> dict[str, list[dict]]:
    """List `<root>/judges/*.{yaml,yml}` for each remote_root.

    Non-recursive — `judges/` is a single per-root folder, same convention
    as `prompts/` (see `_load_prompt_template`'s docstring: `remote_root` is
    always a top-level configured folder, never a nested one).

    Args:
        client: WebDAV client.
        remote_roots: Root remote folders to scan.
        tmp_dir: Local scratch directory for downloading judge YAMLs.

    Returns:
        `{remote_root: [judge_config, ...]}`. A malformed judge YAML logs an
        error and is skipped — it does not abort discovery of other judges.
    """
    result: dict[str, list[dict]] = {}
    for root in remote_roots:
        judges_dir = root.rstrip("/") + f"/{JUDGES_SUBFOLDER}/"
        try:
            entries = client.list(judges_dir, get_info=True)
        except RemoteResourceNotFound:
            continue
        judges = []
        for e in entries:
            if e["isdir"] or Path(e["path"]).suffix.lower() not in SIDECAR_SUFFIXES:
                continue
            name = Path(e["path"]).stem
            try:
                text = _download_text(client, e["path"], tmp_dir)
                judges.append(_parse_judge_yaml(text, name))
            except Exception as exc:
                logger.error("Failed to load judge '%s' in %s: %s", name, judges_dir, exc)
        if judges:
            result[root] = judges
    return result


def _locate_score_token(logprobs_content: list[dict], scale: dict) -> list[dict] | None:
    """Find the `top_logprobs` candidate list at the position holding the
    judge's answer token.

    Reasoning-tuned judge models (e.g. glm-4-7-flash without a configured
    vLLM `--reasoning-parser`) emit `<think>...</think>` inline as literal
    text in `content` — so the answer token is not necessarily the first
    generated token. This walks the token sequence, and once `</think>`
    appears in the concatenated running text, resumes scanning from the
    entry right after the one that completed it (if `</think>` never
    appears at all, scans from index 0). Confirmed live (2026-08-28,
    glm-4-7-flash on vLLM v0.28.0 with `--reasoning-parser=glm45`
    configured): even once a reasoning-parser splits the message-level
    `content`/`reasoning` fields cleanly, `logprobs.content` stays ONE flat
    raw token stream covering reasoning + answer together and still
    contains a literal `</think>` token — so this scan is needed
    unconditionally, not just as a fallback for the no-parser case. From the
    start index, returns the first entry whose own *sampled* token matches a
    valid scale value.

    Args:
        logprobs_content: `choices[0]["logprobs"]["content"]` from a judge
            chat-completion response — a list of
            `{"token", "logprob", "top_logprobs": [{"token","logprob"}, ...]}`.
        scale: Parsed scale, as returned by `_parse_judge_scale`.

    Returns:
        The matching entry's `top_logprobs` list, or `None` if no token in
        the sequence matches any valid scale value.
    """
    running, start = "", 0
    for i, entry in enumerate(logprobs_content):
        running += entry.get("token", "")
        if "</think>" in running:
            start = i + 1
            break

    valid = {str(v).lower() for v in scale["values"]}
    for entry in logprobs_content[start:]:
        if entry.get("token", "").strip().lower() in valid:
            return entry.get("top_logprobs") or None
    return None


def _score_from_logprobs(top_logprobs: list[dict], scale: dict) -> float | None:
    """Probability-weighted G-Eval score over the candidate tokens at one position.

    For a numeric scale, the candidate value is the number itself. For a
    label scale, the candidate value is the 0-based ordinal index in
    `scale["values"]` (ascending-quality order) — this is what makes
    `kind: logprob` produce a continuous score even for a discrete label
    scale.

    `weight(v) = exp(logprob)` for each `top_logprobs` entry whose token
    (stripped, case-insensitive) matches candidate value `v` (summed if the
    same value appears via more than one token variant, e.g. `"4"` and
    `" 4"`). `score = sum(v * weight(v)) / sum(weight(v))`, over matched `v`
    only.

    Args:
        top_logprobs: The candidate list from `_locate_score_token`.
        scale: Parsed scale, as returned by `_parse_judge_scale`.

    Returns:
        The weighted-average score, or `None` if zero entries match any
        scale value.
    """
    if scale["type"] == "numeric":
        candidates = {str(v).lower(): float(v) for v in scale["values"]}
    else:
        candidates = {v.lower(): float(i) for i, v in enumerate(scale["values"])}

    weights: dict[float, float] = {}
    for entry in top_logprobs:
        tok = str(entry.get("token", "")).strip().lower()
        if tok in candidates:
            v = candidates[tok]
            weights[v] = weights.get(v, 0.0) + math.exp(entry["logprob"])

    total = sum(weights.values())
    if total == 0:
        return None
    return sum(v * w for v, w in weights.items()) / total


def _score_from_label(text: str, scale: dict) -> float | str | None:
    """Parse a plain-text judge reply against the scale (`kind: label`).

    Numeric: strips whitespace/trailing punctuation and tries `int()`/
    `float()`; on failure, regex-searches for the first standalone integer
    substring that's a valid scale value. Label: exact match (stripped,
    case-insensitive) against `scale["values"]`; if no exact match, falls
    back to "exactly one label appears as a substring of the reply".

    Args:
        text: The judge LLM's raw reply text.
        scale: Parsed scale, as returned by `_parse_judge_scale`.

    Returns:
        The matched value (float for numeric, the original-cased label
        string for label scales), or `None` if parsing failed.
    """
    cleaned = text.strip().rstrip(".!")
    if scale["type"] == "numeric":
        try:
            return float(cleaned)
        except ValueError:
            pass
        valid = {str(v) for v in scale["values"]}
        for m in re.finditer(r"-?\d+(?:\.\d+)?", cleaned):
            if m.group(0) in valid:
                return float(m.group(0))
        return None

    for value in scale["values"]:
        if cleaned.lower() == value.lower():
            return value
    matches = [v for v in scale["values"] if v.lower() in cleaned.lower()]
    return matches[0] if len(matches) == 1 else None


async def _call_judge_llm(
    session: aiohttp.ClientSession, model: str, prompt: str, want_logprobs: bool
) -> tuple[str | None, list[dict] | None]:
    """Send a judge chat-completion request to the in-cluster KubeAI endpoint.

    Same endpoint/headers/error-handling/`</think>`-stripping shape as
    `_call_llm` (duplicated deliberately — `_call_llm` must not change for
    the non-judge path). No `max_tokens` cap is set (see `JUDGE_TOP_LOGPROBS`'s
    docstring note above: a reasoning-tuned model must be allowed to finish
    its `<think>` block before it ever reaches the answer).

    Args:
        session: Shared aiohttp session.
        model: KubeAI model name to use as judge.
        prompt: Fully rendered judge prompt text.
        want_logprobs: If True, requests `logprobs`/`top_logprobs` in the
            payload (for `kind: logprob`).

    Returns:
        `(content, logprobs_content)`:
        - content: stripped reply text, or `None` on any HTTP/network
          failure (identical contract to `_call_llm`'s return).
        - logprobs_content: the raw `choices[0]["logprobs"]["content"]` list
          when `want_logprobs` was True AND the server returned a
          well-formed non-empty logprobs field, else `None`. A `None` here
          (not an empty list, not an exception) is the signal that the
          server did not give us usable logprobs — the graceful-degradation
          contract for this cluster's unverified KubeAI/vLLM logprobs
          support.
    """
    llm_url = os.environ.get("LLM_URL", "http://kubeai.llm.svc.cluster.local/openai/v1")
    endpoint = f"{llm_url.rstrip('/')}/chat/completions"
    payload = {"model": model, "messages": [{"role": "user", "content": prompt}]}
    if want_logprobs:
        payload["logprobs"] = True
        payload["top_logprobs"] = JUDGE_TOP_LOGPROBS
    headers = {"Authorization": "Bearer not-used"}

    try:
        async with session.post(endpoint, json=payload, headers=headers) as resp:
            if not (200 <= resp.status < 400):
                body = await resp.text()
                logger.error("Judge LLM request failed (HTTP %s): %s", resp.status, body)
                return None, None
            data = json.loads(await resp.read())
            content = data["choices"][0]["message"]["content"]
            if "</think>" in content:
                content = content.rsplit("</think>", 1)[1]
            content = content.strip()

            logprobs_content = None
            if want_logprobs:
                try:
                    logprobs_content = data["choices"][0]["logprobs"]["content"] or None
                except (KeyError, TypeError):
                    logprobs_content = None
            return content, logprobs_content
    except Exception as exc:
        logger.error("Error calling judge LLM: %s: %s", type(exc).__name__, exc)
        return None, None


def _read_judge_scores_xlsx(local_path: Path) -> dict[str, dict[str, dict[str, float | str]]]:
    """Parse a downloaded `<judge>_scores.xlsx` into `{prompt: {stem: {model: score}}}`.

    Each jurisdiction prompt gets its own sheet, named after the prompt
    (truncated to Excel's 31-char sheet-name limit, and subject to
    openpyxl's own sanitizing/dedup of invalid characters — a strong
    assumption that jurisdiction prompt names are short simple slugs, true
    of every prompt name in this codebase's actual use so far). The
    "Overview" sheet (a derived cross-prompt comparison, see
    `_write_judge_scores_xlsx`) is skipped — it's not raw per-cell data to
    resume scoring from. Blank cells (unscored) are skipped, never treated
    as a score of 0/None.

    A score can be numeric (`kind: logprob`, or `kind: label` on a numeric
    scale) or a label string (`kind: label` on a non-numeric scale, e.g.
    poor/fair/good/excellent — see `_score_from_label`). openpyxl already
    returns each cell as its native stored type, so no conversion is
    needed — just pass the value through as-is (forcing `float()`
    unconditionally here previously crashed on any label-string score).

    Args:
        local_path: Local path of a downloaded scores workbook.

    Returns:
        `{prompt: {stem: {model: score}}}`.
    """
    from openpyxl import load_workbook

    wb = load_workbook(local_path, read_only=True, data_only=True)
    result: dict[str, dict[str, dict[str, float | str]]] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name == "Overview":
            continue
        rows = wb[sheet_name].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        models = list(header[1:])

        prompt_scores: dict[str, dict[str, float | str]] = {}
        for row in rows:
            if not row or row[0] is None:
                continue
            stem = str(row[0])
            for model, value in zip(models, row[1:]):
                if value is not None:
                    prompt_scores.setdefault(stem, {})[str(model)] = value
        if prompt_scores:
            result[sheet_name] = prompt_scores
    return result


def _load_existing_scores(
    client: Client, remote_path: str, tmp_dir: Path
) -> dict[str, dict[str, dict[str, float | str]]]:
    """Download+parse an existing scores workbook, or `{}` if it doesn't exist yet.

    Args:
        client: WebDAV client.
        remote_path: Remote path of the `<judge>_scores.xlsx`.
        tmp_dir: Local scratch directory for the temporary download.

    Returns:
        `{prompt: {stem: {model: score}}}`, `{}` if the file doesn't exist remotely.
    """
    local_path = tmp_dir / Path(remote_path).name
    try:
        client.download_sync(remote_path=remote_path, local_path=str(local_path))
    except RemoteResourceNotFound:
        return {}
    try:
        return _read_judge_scores_xlsx(local_path)
    finally:
        local_path.unlink(missing_ok=True)


def _batch_row_for_transcript_stem(
    transcript_stem: str, batch_rows: dict[str, dict]
) -> dict | None:
    """Find the batch row whose file produced a given transcript stem.

    `batch_rows` is keyed by the original audio file's stem, but a
    transcript stem can be `<stem>` (no `stt:` override) or `<stem>_<model>`
    (one per listed STT model) — the same expected-stems derivation
    `_scan_prompt_cells`/`_collect_llm_jobs` use, needed here to find the
    original row's `context:` field for a given already-scanned cell.

    Args:
        transcript_stem: A transcript stem as it appears in `transcriptions/`.
        batch_rows: `{stem: row_dict}`, as returned by `_load_batch_rows`.

    Returns:
        The matching row dict, or `None` if no row produces this stem.
    """
    for stem, data in batch_rows.items():
        stt_models = _stt_models_from_dict(data)
        expected = [stem] if not stt_models else [f"{stem}_{m}" for m in stt_models]
        if transcript_stem in expected:
            return data
    return None


def _missing_judge_cells(
    cells: dict[str, dict[str, str]], existing_scores: dict[str, dict[str, float]]
) -> list[tuple[str, str, str]]:
    """Pure set-difference: cells needing a score.

    This is the entire "don't re-score" guarantee — isolated as its own
    pure, I/O-free function specifically so it can be unit-tested in
    isolation.

    Args:
        cells: `{stem: {model: output_filename}}` — one prompt's worth, from
            `_scan_prompt_cells`.
        existing_scores: `{stem: {model: score}}`, from `_load_existing_scores`.

    Returns:
        List of `(stem, model, output_filename)` triples present in `cells`
        but absent from `existing_scores`.
    """
    return [
        (stem, model, filename)
        for stem, models in cells.items()
        for model, filename in models.items()
        if model not in existing_scores.get(stem, {})
    ]


def _collect_judge_scoring_groups(
    client: Client,
    batch_folders: dict[str, dict[str, dict]],
    folder_roots: dict[str, str],
    judges_by_root: dict[str, list[dict]],
    tmp_dir: Path,
) -> list[dict]:
    """Build one scoring group per (batch folder, judge) with at least one
    jurisdiction prompt having existing LLM output cells.

    One group == one output workbook (`llm/<judge>_scores.xlsx`) — a judge
    scoring multiple jurisdiction prompts in the same folder gets ONE file
    with one sheet per prompt, not one file per prompt, so a person can
    compare a judge's take across different prompts side by side instead
    of hunting across scattered files.

    Args:
        client: WebDAV client.
        batch_folders: `{folder_path: batch_rows}`, from `_collect_llm_jobs`.
        folder_roots: `{folder_path: top_level_root}`, from `_collect_llm_jobs`.
        judges_by_root: `{remote_root: [judge_config, ...]}`, from `_collect_judges`.
        tmp_dir: Local scratch directory for downloads.

    Returns:
        List of dicts: `{folder_path, llm_dir, judge, batch_rows,
        scores_remote_path, prompts: {prompt: {cells, existing_scores
        (loaded once, mutated in place by the caller as new scores land)}}}`.
    """
    groups = []
    for folder_path, batch_rows in batch_folders.items():
        judges = judges_by_root.get(folder_roots[folder_path], [])
        if not judges:
            continue
        prompt_cells = _scan_prompt_cells(client, folder_path, batch_rows)
        llm_dir = folder_path.rstrip("/") + f"/{LLM_SUBFOLDER}/"
        for judge in judges:
            scores_remote_path = llm_dir + f"{judge['name']}_scores.xlsx"
            existing_by_prompt = _load_existing_scores(client, scores_remote_path, tmp_dir)
            prompts_data: dict[str, dict] = {}
            for prompt in judge["jurisdiction"]:
                cells = prompt_cells.get(prompt)
                if not cells:
                    continue
                prompts_data[prompt] = {
                    "cells": cells,
                    "existing_scores": existing_by_prompt.get(prompt, {}),
                }
            if not prompts_data:
                continue
            groups.append(
                {
                    "folder_path": folder_path,
                    "llm_dir": llm_dir,
                    "judge": judge,
                    "batch_rows": batch_rows,
                    "prompts": prompts_data,
                    "scores_remote_path": scores_remote_path,
                }
            )
    return groups


def _write_judge_scores_xlsx(
    local_path: Path, prompts_data: dict[str, dict[str, dict[str, float | str]]]
) -> None:
    """Write one judge's scores workbook: one sheet per jurisdiction prompt,
    plus an "Overview" sheet comparing prompts against each other.

    Per-prompt sheet (named after the prompt, truncated to Excel's 31-char
    limit): header `["filename"] + models`, one row per stem (sorted),
    numeric cell per (stem, model), blank (`None`, not 0) if unscored —
    blank is what makes `_missing_judge_cells` correctly treat it as
    still-needs-scoring on a future run. A bar chart is embedded on the
    same sheet, built directly off that sheet's own cell range.

    "Overview" (first sheet): one row per prompt, one column per model,
    cell = that (prompt, model)'s average score across all its stems — the
    actual "compare prompts in this judge's jurisdiction" view, since the
    per-prompt sheets alone only let you compare models within one prompt.
    Plus its own bar chart (categories = prompts, one series per model).

    Full rewrite of the file's contents each call (old + new scores already
    merged into each prompt's score map by the caller before this is
    invoked).

    Args:
        local_path: Local path to save the workbook to.
        prompts_data: `{prompt: {stem: {model: score}}}` — every
            jurisdiction prompt this judge has data for in this folder.
    """
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    wb.remove(wb.active)

    overview_rows: list[tuple[str, dict[str, float]]] = []
    all_models: set[str] = set()
    for prompt in sorted(prompts_data):
        score_map = prompts_data[prompt]
        models = sorted({m for row in score_map.values() for m in row})
        all_models.update(models)
        ws = wb.create_sheet(prompt[:31])
        ws.append(["filename"] + models)
        for stem in sorted(score_map):
            ws.append([stem] + [score_map[stem].get(m) for m in models])

        n_rows = len(score_map)
        if models and n_rows:
            chart = BarChart()
            chart.type, chart.title = "col", prompt
            chart.y_axis.title, chart.x_axis.title = "score", "transcript"
            data = Reference(ws, min_col=2, max_col=1 + len(models), min_row=1, max_row=1 + n_rows)
            cats = Reference(ws, min_col=1, min_row=2, max_row=1 + n_rows)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, ws.cell(row=1, column=len(models) + 3).coordinate)

        # kind: label with a non-numeric scale (e.g. poor/fair/good/excellent)
        # stores the label string itself as the score (see _score_from_label) —
        # only numeric scores (numeric-scale label judges, or any kind: logprob
        # judge, which always produces a float) can be averaged. A model whose
        # scores are all non-numeric labels still gets its own Overview
        # column (for a consistent model list across every sheet) but with a
        # blank average, rather than crashing or silently coercing strings to 0.
        averages = {}
        for model in models:
            numeric_values = [
                row[model] for row in score_map.values() if isinstance(row.get(model), (int, float))
            ]
            if numeric_values:
                averages[model] = sum(numeric_values) / len(numeric_values)
        overview_rows.append((prompt, averages))

    overview = wb.create_sheet("Overview", 0)
    sorted_models = sorted(all_models)
    overview.append(["prompt"] + sorted_models)
    for prompt, avg in overview_rows:
        overview.append([prompt] + [avg.get(m) for m in sorted_models])

    n_prompts = len(overview_rows)
    if sorted_models and n_prompts:
        chart = BarChart()
        chart.type, chart.title = "col", "Prompt comparison"
        chart.y_axis.title, chart.x_axis.title = "avg score", "prompt"
        data = Reference(
            overview, min_col=2, max_col=1 + len(sorted_models), min_row=1, max_row=1 + n_prompts
        )
        cats = Reference(overview, min_col=1, min_row=2, max_row=1 + n_prompts)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        overview.add_chart(chart, overview.cell(row=1, column=len(sorted_models) + 3).coordinate)

    wb.save(local_path)


def _upload_judge_group_scores(client: Client, group: dict, tmp_dir: Path) -> None:
    """Write and upload one group's scores.xlsx from its current `existing_scores`.

    Called as soon as a group's cells are all attempted (see `main()`) — not
    batched to the end of the whole judge pass, so a large backlog's results
    become visible incrementally, group by group, instead of only after
    every single cell across every judge has been scored.

    Args:
        client: WebDAV client.
        group: One entry from `_collect_judge_scoring_groups`'s result, with
            each prompt's `existing_scores` already updated with this run's
            new scores.
        tmp_dir: Local scratch directory for the upload.
    """
    prompts_data = {prompt: pdata["existing_scores"] for prompt, pdata in group["prompts"].items()}
    local = tmp_dir / Path(group["scores_remote_path"]).name
    _write_judge_scores_xlsx(local, prompts_data)
    try:
        client.list(group["llm_dir"])
    except RemoteResourceNotFound:
        client.mkdir(group["llm_dir"])
    try:
        client.upload_sync(local_path=str(local), remote_path=group["scores_remote_path"])
        logger.info("Uploaded %s", group["scores_remote_path"])
    except Exception as exc:
        logger.error("Upload failed for %s: %s", group["scores_remote_path"], exc)
    local.unlink(missing_ok=True)


async def _score_one_judge_cell(
    client: Client, session: aiohttp.ClientSession, cell_job: dict, tmp_dir: Path
) -> float | str | None:
    """Score one (stem, model) cell against one judge.

    Downloads the output `.md` being judged, downloads the transcript only
    if the judge's template references `{transcript}` (avoids a needless
    WebDAV round-trip otherwise), renders the judge prompt, and dispatches
    to label or logprob scoring per `judge["kind"]`. Any exception is
    caught, logged, and returns `None` — never propagates (matches
    `_call_llm`'s and the main jobs loop's per-item try/except discipline).

    Args:
        client: WebDAV client.
        session: Shared aiohttp session.
        cell_job: `{**group, "stem", "model", "output_name"}` — one flattened
            work item from `_collect_judge_scoring_groups`'s groups.
        tmp_dir: Local scratch directory for downloads.

    Returns:
        The score (float, or a label string — though `kind: label` on a
        label scale returns the original-cased label, which the caller
        stores as-is), or `None` if scoring failed/was skipped.
    """
    judge = cell_job["judge"]
    try:
        response_text = _download_text(
            client, cell_job["llm_dir"] + cell_job["output_name"], tmp_dir
        )
    except Exception as exc:
        logger.error(
            "Judge '%s': failed to download %s: %s", judge["name"], cell_job["output_name"], exc
        )
        return None

    transcript = ""
    if "{transcript}" in judge["template"]:
        transcript_path = (
            cell_job["folder_path"].rstrip("/") + f"/{TRANSCRIPT_SUBFOLDER}/{cell_job['stem']}.txt"
        )
        try:
            transcript = _download_text(client, transcript_path, tmp_dir)
        except Exception as exc:
            logger.error(
                "Judge '%s': failed to download transcript for %s: %s",
                judge["name"],
                cell_job["stem"],
                exc,
            )

    context = ""
    extra_fields: dict[str, str] = {}
    row_data = _batch_row_for_transcript_stem(cell_job["stem"], cell_job["batch_rows"])
    if row_data is not None:
        parsed_row = _sidecar_from_dict(row_data)
        context = parsed_row["context"]
        extra_fields = parsed_row["extra_fields"]

    prompt = _render_judge_prompt(
        judge["template"], response_text, transcript, context, extra_fields
    )

    if judge["kind"] == "logprob":
        content, logprobs_content = await _call_judge_llm(
            session, judge["llm"], prompt, want_logprobs=True
        )
        if content is None:
            return None
        if not logprobs_content:
            logger.error(
                "Judge '%s': no usable logprobs returned by model %s "
                "(unverified KubeAI/vLLM logprobs support) — skipping cell.",
                judge["name"],
                judge["llm"],
            )
            return None
        top_logprobs = _locate_score_token(logprobs_content, judge["scale"])
        if top_logprobs is None:
            logger.error(
                "Judge '%s': no scale-matching token found in response (scanned %d tokens, "
                "post-</think> aware) — skipping cell.",
                judge["name"],
                len(logprobs_content),
            )
            return None
        score = _score_from_logprobs(top_logprobs, judge["scale"])
        if score is None:
            logger.error(
                "Judge '%s': matched token had no scale-matching alternatives in top_logprobs "
                "(%r) — skipping cell.",
                judge["name"],
                top_logprobs,
            )
        return score

    content, _ = await _call_judge_llm(session, judge["llm"], prompt, want_logprobs=False)
    if content is None:
        return None
    score = _score_from_label(content, judge["scale"])
    if score is None:
        logger.error(
            "Judge '%s': could not parse a valid scale value from reply %r — skipping cell.",
            judge["name"],
            content,
        )
    return score
